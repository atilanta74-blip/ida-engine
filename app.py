import io
import json
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
import google.generativeai as genai
import streamlit as st

# --- OLDALBEÁLLÍTÁSOK ---
st.set_page_config(
    page_title="BSHM / B-ME // IDA",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --- CSICSA & HIGH-TECH STÍLUS ---
st.markdown(
    """
<style>
    .stApp {
        background: linear-gradient(135deg, #0b0f19 0%, #111827 50%, #070a12 100%);
        color: #e2e8f0;
        font-family: 'Segoe UI', Roboto, sans-serif;
    }
    .header-box {
        background: rgba(17, 24, 39, 0.85);
        border: 1px solid #1e293b;
        border-left: 5px solid #0284c7;
        padding: 20px 25px;
        border-radius: 12px;
        box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.5), 0 0 15px rgba(2, 132, 199, 0.2);
        margin-bottom: 25px;
    }
    .header-title {
        font-size: 26px;
        font-weight: 800;
        letter-spacing: 1.5px;
        color: #f8fafc;
        margin: 0;
        text-transform: uppercase;
    }
    .header-sub {
        font-size: 13px;
        color: #38bdf8;
        letter-spacing: 3px;
        font-weight: 600;
        margin-top: 4px;
        text-transform: uppercase;
    }
    .kpi-card {
        background: rgba(30, 41, 59, 0.6);
        border: 1px solid rgba(255, 255, 255, 0.08);
        padding: 14px 18px;
        border-radius: 10px;
        backdrop-filter: blur(10px);
    }
    .kpi-title { font-size: 11px; text-transform: uppercase; letter-spacing: 1px; color: #94a3b8; }
    .kpi-value { font-size: 17px; font-weight: 700; color: #38bdf8; margin-top: 3px; }

    .stTextArea textarea, .stTextInput input {
        background-color: #0f172a !important;
        color: #ffffff !important;
        -webkit-text-fill-color: #ffffff !important;
        font-size: 14px !important;
        line-height: 1.5 !important;
        border: 1px solid #0284c7 !important;
        border-radius: 8px !important;
    }
    .stTextArea textarea:focus, .stTextInput input:focus {
        border-color: #38bdf8 !important;
        box-shadow: 0 0 12px rgba(56, 189, 248, 0.4) !important;
    }
    .stButton>button {
        width: 100%;
        background: linear-gradient(90deg, #0284c7 0%, #2563eb 100%) !important;
        color: #ffffff !important;
        font-size: 16px !important;
        font-weight: 700 !important;
        letter-spacing: 1.5px !important;
        text-transform: uppercase !important;
        border: none !important;
        border-radius: 8px !important;
        padding: 12px 24px !important;
        box-shadow: 0 4px 15px rgba(2, 132, 199, 0.4) !important;
    }
    .stDownloadButton>button {
        width: 100%;
        background: linear-gradient(90deg, #059669 0%, #10b981 100%) !important;
        color: white !important;
        font-weight: 700 !important;
        border-radius: 8px !important;
        letter-spacing: 1px;
        box-shadow: 0 4px 15px rgba(16, 185, 129, 0.4) !important;
    }
</style>
""",
    unsafe_allow_html=True,
)

# --- OLDALSÁV: API KULCS BEÁLLÍTÁSA ---
with st.sidebar:
  st.markdown("### ⚙️ AI Rendszer Beállítás")
  api_key_input = st.text_input(
      "Google Gemini API Kulcs:",
      type="password",
      value=st.secrets.get("GEMINI_API_KEY", ""),
      help="Illeszd be az AI Studio kulcsodat",
  )
  st.markdown("---")
  st.markdown(
      "💡 **Tipp:** Ha a Streamlit Secrets-be beírtad a `GEMINI_API_KEY`-t,"
      " üresen maradhat."
  )

# --- FEJLÉC ÉS KPI STATUS BAR ---
st.markdown(
    """
<div class="header-box">
    <div class="header-title">⚡ BSHM / B-ME // IDA</div>
    <div class="header-sub">POWERED BY NAGY ATTILA</div>
</div>
""",
    unsafe_allow_html=True,
)

col_k1, col_k2, col_k3, col_k4 = st.columns(4)
with col_k1:
  st.markdown(
      '<div class="kpi-card"><div class="kpi-title">Rendszer'
      ' Státusz</div><div class="kpi-value" style="color: #4ade80;">● ONLINE /'
      " AI READY</div></div>",
      unsafe_allow_html=True,
  )
with col_k2:
  st.markdown(
      '<div class="kpi-card"><div class="kpi-title">Sablon Revízió</div><div'
      ' class="kpi-value">Rev. 3 (2019.04.03)</div></div>',
      unsafe_allow_html=True,
  )
with col_k3:
  st.markdown(
      '<div class="kpi-card"><div class="kpi-title">Regisztráció</div><div'
      ' class="kpi-value">BME001-F-BDE-3</div></div>',
      unsafe_allow_html=True,
  )
with col_k4:
  st.markdown(
      '<div class="kpi-card"><div class="kpi-title">Ellenőrzési'
      ' Felelős</div><div class="kpi-value">Janóczki M</div></div>',
      unsafe_allow_html=True,
  )

st.markdown("<br>", unsafe_allow_html=True)

# --- BEVITELI MEZŐK ---
c_left, c_right = st.columns([1, 1.4])
with c_left:
  st.markdown("### 📋 Alapadatok")
  gep_nev = st.text_input("Berendezés / Gép megnevezése", value="RF62")
  hiba_rovid = st.text_input(
      "Hiba megnevezése", value="Rimtető megfogó mh szártörés"
  )
  datum_val = st.text_input("Dátum", value="2026.09.25.")
  felelos_val = st.text_input("Karbantartási Felelős", value="Nagy Attila")

with c_right:
  st.markdown("### 📝 Műszakos Jegyzetek & Tények")
  jegyzet_szoveg = st.text_area(
      "Másold be az eseményt / műszakos jegyzetet:",
      height=188,
      value=(
          "2023.06.07.-én a rf62 gép szétszerelő állomáson a rimtető megfogó"
          " mh egyik szára eltörött. A szakos karbantartás cserélte a mh-t."
          " Ezután ellenőriztem a központosságot és a mérés kimutatta, hogy a 3"
          " irány egyikénél -2mm eltérés van. Ez az eltérés eredményezhette,"
          " hogy a több ezer nyitás-zárás alatt a szár meggyengült és eltörött."
          " Be lett állítva a központosság, majd horizontal check lesz"
          " alkalmazva minden hasonló gépen."
      ),
  )


# --- INTELLIGENS, FLASH MODELLEKET HASZNÁLÓ MOTOR ---
def elemez_geminivel(api_key, gep, hiba, datum, felelos, jegyzet):
  genai.configure(api_key=api_key)

  elerheto = []
  try:
    for m in genai.list_models():
      if "generateContent" in m.supported_generation_methods:
        elerheto.append(m.name)
  except Exception:
    pass

  # KIZÁRJUK a "pro" modelleket, hogy ne fusson 429 kvótahibára
  flash_list = [
      m for m in elerheto if "flash" in m.lower() and "pro" not in m.lower()
  ]

  tartalek_flash = [
      "models/gemini-2.0-flash",
      "models/gemini-1.5-flash-latest",
      "models/gemini-1.5-flash-002",
      "models/gemini-1.5-flash",
      "models/gemini-1.5-flash-8b",
      "gemini-2.0-flash",
      "gemini-1.5-flash",
  ]

  kiprobalando = []
  for m in flash_list + tartalek_flash:
    if m not in kiprobalando:
      kiprobalando.append(m)

  prompt = f"""
    Te egy tapasztalt mechatronikai TPM karbantartó mérnök vagy a Bridgestone gyárban.
    Elemezd az alábbi konkrét műszakos jegyzetet a gyári BSHM B-ME IDA szabvány szerint:

    Berendezés: {gep}
    Hiba: {hiba}
    Dátum: {datum}
    Felelős: {felelos}
    Műszakos jegyzet:
    {jegyzet}

    KIZÁRÓLAG érvényes JSON formátumban válaszolj, az alábbi séma szerint:
    {{
      "kriteriumok": [
        {{
          "nev": "Tisztaság",
          "kerdes": "(A részegység (berendezés) elég tiszta, hogy ellenőrizhető legyen és az állapota ne romoljon rohamosan?)",
          "idealis": "1. ...",
          "aktualis": "1. ...",
          "hianyossag": "1. ... (ha nincs: 1. Nincs hiányosság.)",
          "gyokerok": "1. ... (ha nincs: 1. Nem releváns.)",
          "ellenintezkedes": "1. ... (ha nincs: 1. Standard CIL tisztítás fenntartása.)"
        }},
        ... (mind a 13 kritérium sorrendben: Tisztaság, Meghúzottság, Kenés, Karbantartás, Környezet, Működtetés, Specifikáció, Működés, Telepítés, Összeszerelés, Gyártás, Tervezés, Javítás)
      ],
      "ot_miert_agak": [
        {{
          "ag_nev": "1. ÁG: MŰSZAKI / FIZIKAI ÁG",
          "problema": "Probléma leírás a mechanikai törés / feszültség szempontjából",
          "miert1": "1. Miért kérdés és válasz",
          "miert2": "2. Miért kérdés és válasz",
          "miert3": "3. Miért kérdés és válasz",
          "miert4": "4. Miért kérdés és válasz",
          "miert5_gyokerok": "GYÖKÉROK: A fizikai/mechanikai ok"
        }},
        {{
          "ag_nev": "2. ÁG: KARBANTARTÁSI / MEGELŐZÉSI ÁG",
          "problema": "Probléma leírás a beállítási és megelőzési ellenőrzés szempontjából",
          "miert1": "1. Miért kérdés és válasz",
          "miert2": "2. Miért kérdés és válasz",
          "miert3": "3. Miért kérdés és válasz",
          "miert4": "4. Miért kérdés és válasz",
          "miert5_gyokerok": "GYÖKÉROK: A geometriai / központossági ellenőrzési ciklus hiánya"
        }},
        {{
          "ag_nev": "3. ÁG: HORIZONTÁLIS KITERJESZTÉSI ÁG",
          "problema": "Probléma leírás a hasonló gépekre történő kiterjesztés szempontjából",
          "miert1": "1. Miért kérdés és válasz",
          "miert2": "2. Miért kérdés és válasz",
          "miert3": "3. Miért kérdés és válasz",
          "miert4": "4. Miért kérdés és válasz",
          "miert5_gyokerok": "GYÖKÉROK: A horizontális ellenőrzési standard hiánya"
        }}
      ],
      "vegleges_akcioterv": "Részletes összefoglaló a hibáról, az elvégzett beavatkozásról és a horizontális check bevezetéséről felelőssel és határidővel."
    }}
    Minden mező legyen közvetlenül az adott jegyzetre szabva, szakmai magyar nyelven!
    """

  utolso_hiba = None
  for mod_nev in kiprobalando:
    try:
      model = genai.GenerativeModel(mod_nev)
      try:
        res = model.generate_content(
            prompt, generation_config={"response_mime_type": "application/json"}
        )
        raw_text = res.text.strip()
      except Exception:
        res = model.generate_content(prompt)
        raw_text = res.text.strip()

      if raw_text.startswith("```"):
        lines = raw_text.splitlines()
        if lines[0].startswith("```"):
          lines = lines[1:]
        if lines and lines[-1].startswith("```"):
          lines = lines[:-1]
        raw_text = "\n".join(lines).strip()
        if raw_text.startswith("json"):
          raw_text = raw_text[4:].strip()

      return json.loads(raw_text)
    except Exception as err:
      utolso_hiba = err
      continue

  raise utolso_hiba


# --- DOKUMENTUM GENERÁLÓ MOTOR (100% GYÁRI FORMÁTUM, CSAK JANÓCZKI M!) ---
def general_hivatalos_docx(gep, hiba, datum, felelos, ai_data, jegyzet):
  doc = docx.Document()
  for section in doc.sections:
    section.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

  def set_cell_bg(cell, fill_hex):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    )

  def set_cell_pad(cell, top=35, bottom=35, left=50, right=50):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(
            f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}"'
            f' w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left'
            f' w:w="{left}" w:type="dxa"/><w:right w:w="{right}"'
            ' w:type="dxa"/></w:tcMar>'
        )
    )

  def set_cell_border(cell, sz="4"):
    b_xml = (
        f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}"'
        ' w:space="0" w:color="000000"/><w:left w:val="single"'
        f' w:sz="{sz}" w:space="0" w:color="000000"/><w:bottom w:val="single"'
        f' w:sz="{sz}" w:space="0" w:color="000000"/><w:right w:val="single"'
        f' w:sz="{sz}" w:space="0" w:color="000000"/></w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(parse_xml(b_xml))

  p_title = doc.add_paragraph()
  r = p_title.add_run(
      "Üzemzavar elhárítási napi irányítási rendszer – Berendezés hiányosság"
      " értékelés, mélyanalízis, ellenintézkedési lap"
  )
  r.font.name = "Calibri"
  r.font.size = Pt(11)
  r.font.bold = True

  p_info = doc.add_paragraph()
  p_info.add_run("Üzemzavar leírása: ").bold = True
  p_info.add_run(f"{gep} – {hiba}          ").underline = True
  p_info.add_run("Dátum: ").bold = True
  p_info.add_run(f"{datum}          ").underline = True
  p_info.add_run("Karb. felelős: ").bold = True
  p_info.add_run(f"{felelos}").underline = True

  p1_table = doc.add_table(rows=15, cols=6)
  p1_table.alignment = WD_TABLE_ALIGNMENT.CENTER
  p1_widths = [
      Inches(1.8),
      Inches(1.75),
      Inches(1.75),
      Inches(1.75),
      Inches(1.85),
      Inches(1.8),
  ]

  c_plan = p1_table.cell(0, 0)
  c_plan.merge(p1_table.cell(0, 1)).merge(p1_table.cell(0, 2)).merge(
      p1_table.cell(0, 3)
  )
  c_plan.text = "TERVEZD MEG (PLAN)"
  p1_table.cell(0, 4).text = "VALÓSÍTSD MEG (DO)"
  p1_table.cell(0, 5).text = "ELLENŐRIZD (CHECK), STANDARDIZÁLD (ACT)"

  sub_heads = [
      "KRITÉRIUM",
      "IDEÁLIS ÁLLAPOT",
      "AKTUÁLIS ÁLLAPOT",
      "HIÁNYOSSÁG",
      "Hiányosság gyökéroka\n(ismert, vagy “5 miért” elemzés alapján)",
      "Ellenintézkedés\n(határidő, felelős, standardizálás)",
  ]
  for c_i, h in enumerate(sub_heads):
    p1_table.cell(1, c_i).text = h

  for r_i in range(2):
    for cell in p1_table.rows[r_i].cells:
      set_cell_bg(cell, "F2F2F2")
      set_cell_pad(cell, 50, 50, 50, 50)
      set_cell_border(cell, sz="6")
      p = cell.paragraphs[0]
      p.alignment = WD_ALIGN_PARAGRAPH.CENTER
      if p.runs:
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(7.5)
        p.runs[0].font.bold = True

  kriteriumok = ai_data.get("kriteriumok", [])
  for r_idx, krit in enumerate(kriteriumok, start=2):
    if r_idx >= 15:
      break
    col_vals = [
        f"{krit.get('nev', '')}\n{krit.get('kerdes', '')}",
        krit.get("idealis", "1. Standard állapot."),
        krit.get("aktualis", "1. Megfelelő."),
        krit.get("hianyossag", "1. Nincs hiányosság."),
        krit.get("gyokerok", "1. Nem releváns."),
        krit.get("ellenintezkedes", "1. Standard állapot fenntartása."),
    ]
    for c_i, val in enumerate(col_vals):
      cell = p1_table.cell(r_idx, c_i)
      cell.text = val
      set_cell_bg(cell, "FFFFFF")
      set_cell_pad(cell, 35, 35, 50, 50)
      set_cell_border(cell, sz="4")
      p = cell.paragraphs[0]
      p.paragraph_format.line_spacing = Pt(9)
      if p.runs:
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.size = Pt(7)
        p.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        if c_i == 0:
          p.runs[0].font.bold = True

  for row in p1_table.rows:
    for c_i, w in enumerate(p1_widths):
      row.cells[c_i].width = w

  doc.add_page_break()
  p2_title = doc.add_paragraph()
  r = p2_title.add_run(
      "Üzemzavar elhárítási napi irányítási rendszer – Berendezés hiányosság"
      " értékelés, mélyanalízis, ellenintézkedési lap"
  )
  r.font.name = "Calibri"
  r.font.size = Pt(11)
  r.font.bold = True

  p2_info = doc.add_paragraph()
  p2_info.add_run(
      f"Mélyanalízis (5 Miért elemzés)          Gép: {gep}          Dátum:"
      f" {datum}          Karb. felelős: {felelos}"
  ).bold = True

  p2_table = doc.add_table(rows=12, cols=5)
  p2_table.alignment = WD_TABLE_ALIGNMENT.CENTER
  col_w = Inches(10.7 / 5)

  agak = ai_data.get("ot_miert_agak", [])
  p_defs = [
      agak[i].get("problema", "") if i < len(agak) else "" for i in range(5)
  ]
  whys = [
      [agak[i].get("miert1", "") if i < len(agak) else "" for i in range(5)],
      [agak[i].get("miert2", "") if i < len(agak) else "" for i in range(5)],
      [agak[i].get("miert3", "") if i < len(agak) else "" for i in range(5)],
      [agak[i].get("miert4", "") if i < len(agak) else "" for i in range(5)],
      [
          agak[i].get("miert5_gyokerok", "") if i < len(agak) else ""
          for i in range(5)
      ],
  ]

  for c_i in range(5):
    cell = p2_table.cell(0, c_i)
    cell.text = "PROBLÉMA MEGHATÁROZÁSA"
    set_cell_bg(cell, "F2F2F2")
    set_cell_pad(cell, 35, 35, 40, 40)
    set_cell_border(cell, "6")
    cell.paragraphs[0].runs[0].font.name = "Calibri"
    cell.paragraphs[0].runs[0].font.size = Pt(8)
    cell.paragraphs[0].runs[0].font.bold = True

    cell_p = p2_table.cell(1, c_i)
    cell_p.text = p_defs[c_i]
    set_cell_bg(cell_p, "FFFFFF")
    set_cell_pad(cell_p, 35, 35, 40, 40)
    set_cell_border(cell_p, "4")
    if cell_p.paragraphs[0].runs:
      cell_p.paragraphs[0].runs[0].font.name = "Calibri"
      cell_p.paragraphs[0].runs[0].font.size = Pt(7.5)

  for step in range(5):
    h_row = 2 + step * 2
    d_row = 3 + step * 2
    for c_i in range(5):
      c_h = p2_table.cell(h_row, c_i)
      c_h.text = "( MIÉRT?"
      set_cell_bg(c_h, "F2F2F2")
      set_cell_pad(c_h, 25, 25, 40, 40)
      set_cell_border(c_h, "6")
      c_h.paragraphs[0].runs[0].font.name = "Calibri"
      c_h.paragraphs[0].runs[0].font.size = Pt(7.5)
      c_h.paragraphs[0].runs[0].font.bold = True

      c_d = p2_table.cell(d_row, c_i)
      c_d.text = whys[step][c_i]
      set_cell_bg(c_d, "FFFFFF")
      set_cell_pad(c_d, 35, 35, 40, 40)
      set_cell_border(c_d, "4")
      if c_d.paragraphs[0].runs:
        c_d.paragraphs[0].runs[0].font.name = "Calibri"
        c_d.paragraphs[0].runs[0].font.size = Pt(7.5)
        if step == 4 and c_i < len(agak):
          c_d.paragraphs[0].runs[0].font.bold = True

  for row in p2_table.rows:
    for c_i in range(5):
      row.cells[c_i].width = col_w

  doc.add_paragraph().paragraph_format.space_before = Pt(4)
  p_n = doc.add_paragraph()
  p_n.add_run("JEGYZETEK / AKCIÓTERV:").bold = True
  n_box = doc.add_table(rows=1, cols=1)
  n_box.alignment = WD_TABLE_ALIGNMENT.CENTER
  n_cell = n_box.cell(0, 0)
  n_cell.width = Inches(10.7)
  set_cell_bg(n_cell, "FFFFFF")
  set_cell_pad(n_cell, 50, 50, 60, 60)
  set_cell_border(n_cell, "6")
  n_cell.text = ai_data.get("vegleges_akcioterv") or jegyzet
  n_cell.paragraphs[0].runs[0].font.name = "Calibri"
  n_cell.paragraphs[0].runs[0].font.size = Pt(8.5)

  for section in doc.sections:
    footer = section.footer
    foot_table = footer.add_table(rows=2, cols=6, width=Inches(10.7))
    foot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_f = [
        "Szervezet / Részleg",
        "Revízió",
        "1. kiadás dátuma",
        "Felülvizsgálat dátuma:",
        "Készítő / Felülvizsgáló:",
        "Regisztrációs szám",
    ]
    vals_f = [
        "BSHM / B-ME",
        "3.",
        "2018.07.16",
        "2019.04.03.",
        "Janóczki M",
        "BME001-F-BDE-3",
    ]
    widths_f = [
        Inches(1.8),
        Inches(1.2),
        Inches(1.6),
        Inches(1.8),
        Inches(2.5),
        Inches(1.8),
    ]
    for c_i in range(6):
      c1 = foot_table.cell(0, c_i)
      c1.text = headers_f[c_i]
      set_cell_bg(c1, "F2F2F2")
      set_cell_pad(c1, 20, 20, 30, 30)
      set_cell_border(c1, "4")
      c1.paragraphs[0].runs[0].font.name = "Calibri"
      c1.paragraphs[0].runs[0].font.size = Pt(7)
      c1.paragraphs[0].runs[0].font.bold = True

      c2 = foot_table.cell(1, c_i)
      c2.text = vals_f[c_i]
      set_cell_bg(c2, "FFFFFF")
      set_cell_pad(c2, 20, 20, 30, 30)
      set_cell_border(c2, "4")
      c2.paragraphs[0].runs[0].font.name = "Calibri"
      c2.paragraphs[0].runs[0].font.size = Pt(7.5)
    for r_f in foot_table.rows:
      for c_i, w in enumerate(widths_f):
        r_f.cells[c_i].width = w

  output = io.BytesIO()
  doc.save(output)
  output.seek(0)
  return output


# --- GOMB ÉS FOLYAMAT ---
if st.button("⚡ INTELLIGENS ELEMZÉS & DOKUMENTUM GENERÁLÁSA"):
  api_key = api_key_input.strip()
  if not api_key:
    st.error(
        "⚠️ Kérlek, add meg a Google Gemini API kulcsot a bal oldali sávban!"
    )
  else:
    with st.spinner(
        "🧠 A mesterséges intelligencia elemzi a jegyzetet és felépíti az 5-Miért"
        " ágakat..."
    ):
      try:
        ai_eredmeny = elemez_geminivel(
            api_key, gep_nev, hiba_rovid, datum_val, felelos_val, jegyzet_szoveg
        )
        docx_file = general_hivatalos_docx(
            gep_nev,
            hiba_rovid,
            datum_val,
            felelos_val,
            ai_eredmeny,
            jegyzet_szoveg,
        )

        st.success(
            "✅ Az új hibajelenség mélyanalízise és a dokumentum sikeresen"
            " elkészült!"
        )
        st.download_button(
            label="📥 HIVATALOS WORD DOKUMENTUM LETÖLTÉSE (.DOCX)",
            data=docx_file,
            file_name=(
                f"{gep_nev.replace(' ', '_')}_{hiba_rovid.replace(' ', '_')}_IDA.docx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
        )
      except Exception as e:
        st.error(f"Hiba történt az elemzés során: {e}")
